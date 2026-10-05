"""File input (bulk CSV/Excel upload): each row is one complete, independent tick -- no
state carry-over between rows. Additive only -- reuses normalize_scenario, fleet specs, and
dispatcher.state_volatility_class exactly as they are; nothing here changes a rule,
tolerance, the prompt, or any existing schema field.

Flow: parse_rows -> row_to_scenario (per row) -> normalize_scenario (existing, unchanged) ->
validate_rows wraps both and is what /file-input/preview calls. Validated scenarios are
stored in _RUNS, keyed by run_id and row index, so /file-input/run-row can look one up by
index alone -- an unvalidated or hand-crafted scenario can never be executed.
"""
from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass, field
from typing import Any, Optional

import openpyxl

from app.agents.scenario_agent import normalize_scenario
from app.data import fleet
from app.data.dispatcher import state_volatility_class
from app.data.tuning import FLOOR_BANDS
from app.models.schemas import Battery, Difficulty, EnvironmentState, Objective

MAX_ROWS = 200
MAX_FILE_BYTES = 5 * 1024 * 1024  # 5MB

# Pre-run time estimate: mean latency_ms across every "ok" recorded call in recordings/
# (196 calls at the time this was measured, 2026-10-05) -- a measured figure, not a guess.
# One call per valid row (sequential, per the frontend's run loop); live-mode runs can be
# slower (no cache hits), recorded re-runs of an already-completed row are ~free -- this
# estimates the first-run, no-cache case, which is the one worth warning about up front.
ESTIMATED_SECONDS_PER_CALL = 3.3


class FileInputError(Exception):
    """File-level problem (too big, too many rows, unreadable) -- rejected before any
    row-level validation is attempted."""


@dataclass
class ColumnSpec:
    name: str
    kind: str  # "float" | "int" | "bool" | "str"
    required: bool
    default: Any = None
    unit: str = ""
    description: str = ""


# Single source of truth for the template, the data dictionary, and parsing/coercion.
TEMPLATE_COLUMNS: list[ColumnSpec] = [
    ColumnSpec("tick", "int", False, None, "", "Row label only -- defaults to the row's position (1-based) if blank. Never affects scoring."),
    ColumnSpec("seed", "int", False, None, "", "Record-keeping only -- defaults to a value derived from the row position if blank."),
    ColumnSpec("objective", "str", False, "", "", "One of cost_efficiency / min_carbon / max_renewable_utilisation / max_profit, or blank for cost_efficiency."),
    ColumnSpec("solar_output_mw", "float", True, None, "MW", "Current solar generation."),
    ColumnSpec("solar_forecast_mw", "float", True, None, "MW", "Forecast solar generation (informational)."),
    ColumnSpec("wind_output_mw", "float", True, None, "MW", "Current wind generation."),
    ColumnSpec("wind_forecast_mw", "float", True, None, "MW", "Forecast wind generation (informational)."),
    ColumnSpec("base_demand_mw", "float", True, None, "MW", "Demand excluding industrial load."),
    ColumnSpec("base_demand_forecast_mw", "float", True, None, "MW", "Forecast base demand (informational)."),
    ColumnSpec("industrial_demand_mw", "float", True, None, "MW", "Industrial load, added to base_demand_mw to get total demand (computed automatically)."),
    ColumnSpec("grid_frequency_hz", "float", True, None, "Hz", "Grid frequency this tick."),
    ColumnSpec("transmission_constraint_mw", "float", True, None, "MW", "Fixed interconnection line rating."),
    ColumnSpec("transmission_headroom_mw", "float", True, None, "MW", "Available net-flow capacity this tick (<= transmission_constraint_mw)."),
    ColumnSpec("battery_1_soc_pct", "float", True, None, "%", "battery_1's state of charge this tick, 0-100. This row's own value -- never carried over from another row."),
    ColumnSpec("battery_1_available", "bool", False, True, "", "Whether battery_1 is online this tick."),
    ColumnSpec("battery_2_soc_pct", "float", True, None, "%", "battery_2's state of charge this tick, 0-100."),
    ColumnSpec("battery_2_available", "bool", False, True, "", "Whether battery_2 is online this tick."),
    ColumnSpec("previous_floor_pct", "float", False, 25.0, "%", "Reserve floor applied on the previous tick (for the ramp-down rule)."),
    ColumnSpec("electricity_price_per_mwh", "float", True, None, "$/MWh", "Quoted price; buy/sell prices are computed automatically from it."),
    ColumnSpec("carbon_price_per_ton", "float", True, None, "$/t", "Informational only -- never added to any cost metric."),
    ColumnSpec("grid_carbon_intensity_t_per_mwh", "float", True, None, "t/MWh", "Emissions per MWh of grid import."),
    ColumnSpec("demand_response_incentive_per_mwh", "float", True, None, "$/MWh", "Incentive rate if demand response is triggered."),
    ColumnSpec("weather_forecast", "str", False, "n/a", "", "Free-text, informational only."),
    ColumnSpec("storm_alert", "bool", False, False, "", "Whether a storm alert is active this tick."),
    ColumnSpec("maintenance_scheduled", "bool", False, False, "", "Whether maintenance is scheduled this tick."),
    ColumnSpec("events", "str", False, "", "", "Comma-separated event tags, e.g. battery_2_offline,transmission_at_capacity. Blank for none."),
]
_COLUMNS_BY_NAME = {c.name: c for c in TEMPLATE_COLUMNS}


# ---- parsing ---------------------------------------------------------------------------


def parse_rows(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    if len(file_bytes) > MAX_FILE_BYTES:
        raise FileInputError(f"File is {len(file_bytes) / 1024 / 1024:.1f}MB, over the {MAX_FILE_BYTES / 1024 / 1024:.0f}MB limit.")

    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = _parse_csv(file_bytes)
    elif lower.endswith(".xlsx"):
        rows = _parse_xlsx(file_bytes)
    else:
        raise FileInputError(f"Unsupported file type '{filename}' -- upload a .csv or .xlsx file.")

    if len(rows) > MAX_ROWS:
        raise FileInputError(f"File has {len(rows)} rows, over the {MAX_ROWS}-row limit.")
    if not rows:
        raise FileInputError("File has no data rows.")
    return rows


def _parse_csv(file_bytes: bytes) -> list[dict[str, Any]]:
    text = file_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def _parse_xlsx(file_bytes: bytes) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb["Scenario Rows"] if "Scenario Rows" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    rows = []
    for raw in rows_iter:
        if all(v is None for v in raw):
            continue
        rows.append({header[i]: raw[i] for i in range(len(header)) if i < len(raw)})
    return rows


# ---- coercion and row -> scenario -------------------------------------------------------


def _coerce(col: ColumnSpec, raw: Any) -> Any:
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        if col.required:
            raise ValueError(f"'{col.name}' is required.")
        return col.default
    if col.kind == "str":
        return str(raw).strip()
    if col.kind == "bool":
        if isinstance(raw, bool):
            return raw
        s = str(raw).strip().lower()
        if s in ("true", "1", "yes", "y"):
            return True
        if s in ("false", "0", "no", "n"):
            return False
        raise ValueError(f"'{col.name}' should be true/false, got '{raw}'.")
    try:
        value = float(raw)
    except (TypeError, ValueError):
        raise ValueError(f"'{col.name}' should be a number, got '{raw}'.")
    if col.kind == "int":
        return int(value)
    return value


@dataclass
class RowResult:
    row_index: int
    status: str  # "ok" | "error"
    tick: Optional[int] = None
    scenario: Optional[EnvironmentState] = None
    errors: list[str] = field(default_factory=list)
    volatility_class: Optional[str] = None
    volatility_method: Optional[str] = None
    expected_floor_band: Optional[str] = None


def row_to_scenario(row_dict: dict[str, Any], row_index: int) -> EnvironmentState:
    """Raises ValueError with every problem found (not just the first) joined by '; '."""
    errors: list[str] = []
    values: dict[str, Any] = {}
    for col in TEMPLATE_COLUMNS:
        raw = row_dict.get(col.name)
        try:
            values[col.name] = _coerce(col, raw)
        except ValueError as e:
            errors.append(str(e))

    if errors:
        raise ValueError("; ".join(errors))

    # semantic checks beyond type coercion
    for soc_field in ("battery_1_soc_pct", "battery_2_soc_pct"):
        v = values[soc_field]
        if not (0.0 <= v <= 100.0):
            errors.append(f"'{soc_field}' must be between 0 and 100, got {v}.")
    if values["transmission_headroom_mw"] > values["transmission_constraint_mw"]:
        errors.append("'transmission_headroom_mw' cannot exceed 'transmission_constraint_mw'.")
    for mw_field in ("solar_output_mw", "wind_output_mw", "base_demand_mw", "industrial_demand_mw", "transmission_constraint_mw", "transmission_headroom_mw"):
        if values[mw_field] < 0:
            errors.append(f"'{mw_field}' cannot be negative, got {values[mw_field]}.")
    for price_field in ("electricity_price_per_mwh", "carbon_price_per_ton", "grid_carbon_intensity_t_per_mwh", "demand_response_incentive_per_mwh"):
        if values[price_field] < 0:
            errors.append(f"'{price_field}' cannot be negative, got {values[price_field]}.")

    objective_raw = (values.get("objective") or "").strip().lower()
    objective = None
    if objective_raw:
        try:
            objective = Objective(objective_raw)
        except ValueError:
            errors.append(f"'objective' must be one of {[o.value for o in Objective]} or blank, got '{objective_raw}'.")

    if errors:
        raise ValueError("; ".join(errors))

    events = [e.strip() for e in (values.get("events") or "").split(",") if e.strip()]
    tick = values.get("tick") if values.get("tick") is not None else row_index + 1
    seed = values.get("seed") if values.get("seed") is not None else 1_000_000 + row_index

    battery_specs_by_id = {spec["id"]: spec for spec in fleet.BATTERY_SPECS}
    batteries = [
        Battery(
            id="battery_1", min_safe_soc_pct=10.0,
            state_of_charge_pct=values["battery_1_soc_pct"], available=values["battery_1_available"],
            **{k: v for k, v in battery_specs_by_id["battery_1"].items() if k != "id"},
        ),
        Battery(
            id="battery_2", min_safe_soc_pct=10.0,
            state_of_charge_pct=values["battery_2_soc_pct"], available=values["battery_2_available"],
            **{k: v for k, v in battery_specs_by_id["battery_2"].items() if k != "id"},
        ),
    ]

    total_demand = round(values["base_demand_mw"] + values["industrial_demand_mw"], 1)
    total_demand_forecast = round(values["base_demand_forecast_mw"] + values["industrial_demand_mw"], 1)

    from app.data import physics
    buy_price = physics.buy_price_per_mwh(values["electricity_price_per_mwh"])
    sell_price = physics.sell_price_per_mwh(values["electricity_price_per_mwh"])

    # Two-pass floor-band derivation: state_volatility_class never reads expected_floor_*,
    # so a placeholder is safe here -- it's replaced with the real derived band below.
    draft = EnvironmentState(
        tick=tick, seed=seed, difficulty=Difficulty.FILE_INPUT, objective=objective,
        solar_output_mw=values["solar_output_mw"], solar_forecast_mw=values["solar_forecast_mw"],
        wind_output_mw=values["wind_output_mw"], wind_forecast_mw=values["wind_forecast_mw"],
        solar_farms=[], wind_farms=[],
        base_demand_mw=values["base_demand_mw"], base_demand_forecast_mw=values["base_demand_forecast_mw"],
        total_demand_mw=total_demand, total_demand_forecast_mw=total_demand_forecast,
        grid_frequency_hz=values["grid_frequency_hz"],
        transmission_constraint_mw=values["transmission_constraint_mw"], transmission_headroom_mw=values["transmission_headroom_mw"],
        batteries=batteries, previous_floor_pct=values["previous_floor_pct"],
        electricity_price_per_mwh=values["electricity_price_per_mwh"], buy_price_per_mwh=buy_price, sell_price_per_mwh=sell_price,
        carbon_price_per_ton=values["carbon_price_per_ton"], grid_carbon_intensity_t_per_mwh=values["grid_carbon_intensity_t_per_mwh"],
        demand_response_incentive_per_mwh=values["demand_response_incentive_per_mwh"],
        weather_forecast=values["weather_forecast"], storm_alert=values["storm_alert"], maintenance_scheduled=values["maintenance_scheduled"],
        industrial_demand_mw=values["industrial_demand_mw"], events=events,
        expected_behavior="n/a (file input)", expected_floor_min_pct=20.0, expected_floor_max_pct=60.0,
        expected_ladder_step="n/a (file input)", expected_emergency=False,
    )
    vol_class, vol_method = state_volatility_class(draft)
    floor_lo, floor_hi = FLOOR_BANDS[vol_class]
    scenario = draft.model_copy(update={"expected_floor_min_pct": floor_lo, "expected_floor_max_pct": floor_hi})
    return scenario


def validate_rows(parsed_rows: list[dict[str, Any]]) -> list[RowResult]:
    results = []
    for i, row_dict in enumerate(parsed_rows):
        try:
            scenario = row_to_scenario(row_dict, i)
            normalized, _corrections = normalize_scenario(scenario)
            # Idempotent re-derivation for display only (same inputs row_to_scenario already
            # used to pick the floor band) -- avoids stashing non-schema attributes on a
            # pydantic model.
            vol_class, vol_method = state_volatility_class(normalized)
            results.append(RowResult(
                row_index=i, status="ok", tick=normalized.tick, scenario=normalized,
                volatility_class=vol_class, volatility_method=vol_method,
                expected_floor_band=f"{normalized.expected_floor_min_pct}-{normalized.expected_floor_max_pct}%",
            ))
        except ValueError as e:
            results.append(RowResult(row_index=i, status="error", errors=str(e).split("; ")))
        except Exception as e:  # pydantic ValidationError or anything else unexpected
            results.append(RowResult(row_index=i, status="error", errors=[str(e)]))
    return results


# ---- run store (preview -> run-row) ------------------------------------------------------

_RUNS: dict[str, list[Optional[EnvironmentState]]] = {}


def run_id_for(file_bytes: bytes) -> str:
    return "fileinput_" + hashlib.sha256(file_bytes).hexdigest()[:12]


def store_run(run_id: str, results: list[RowResult]) -> None:
    _RUNS[run_id] = [r.scenario if r.status == "ok" else None for r in results]


def get_validated_scenario(run_id: str, row_index: int) -> Optional[EnvironmentState]:
    scenarios = _RUNS.get(run_id)
    if scenarios is None or row_index < 0 or row_index >= len(scenarios):
        return None
    return scenarios[row_index]


# ---- template / example file generation --------------------------------------------------


def _write_rows_sheet(wb: openpyxl.Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "Scenario Rows"
    headers = [c.name for c in TEMPLATE_COLUMNS]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def _write_dictionary_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Data Dictionary")
    ws.append(["Column", "Type", "Required", "Unit", "Description"])
    for col in TEMPLATE_COLUMNS:
        ws.append([col.name, col.kind, "yes" if col.required else "no", col.unit, col.description])


def make_template_file() -> bytes:
    wb = openpyxl.Workbook()
    example_row = {
        "tick": 1, "seed": 1, "objective": "", "solar_output_mw": 40.0, "solar_forecast_mw": 40.0,
        "wind_output_mw": 61.7, "wind_forecast_mw": 61.7, "base_demand_mw": 100.0, "base_demand_forecast_mw": 100.0,
        "industrial_demand_mw": 39.0, "grid_frequency_hz": 50.0, "transmission_constraint_mw": 180.0, "transmission_headroom_mw": 180.0,
        "battery_1_soc_pct": 50.0, "battery_1_available": True, "battery_2_soc_pct": 50.0, "battery_2_available": True,
        "previous_floor_pct": 25.0, "electricity_price_per_mwh": 60.0, "carbon_price_per_ton": 30.0,
        "grid_carbon_intensity_t_per_mwh": 0.6, "demand_response_incentive_per_mwh": 20.0,
        "weather_forecast": "clear", "storm_alert": False, "maintenance_scheduled": False, "events": "",
    }
    _write_rows_sheet(wb, [example_row])
    _write_dictionary_sheet(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EXAMPLE_DAY_ROWS: list[dict[str, Any]] = [
    # A realistic 12-row day, every 2 hours. Independently specified per row (no state
    # carry-over is enforced by the system) but written to look like a coherent day.
    # Deliberately exercises the system rather than one state (2026-10-05 rebuild):
    # - 3 genuine surplus ticks (5, 6, 7) with export headroom DELIBERATELY reduced
    #   (transmission_headroom_mw=12, not the usual 180) and total battery charge headroom
    #   capped at 16MW -- min_required_curtailment_mw is nonzero on all three (7/16/12MW), so
    #   sell, charge, AND curtail are all genuinely in play, not just theoretically available.
    # - cost-vs-carbon can disagree twice: tick 4 is cheap AND dirty (price 28, intensity
    #   0.85, both during a real shortfall so the import-vs-discharge choice matters);
    #   tick 9 is expensive AND clean (price 190, intensity 0.32), the mirror case.
    # - tick 2: battery_2 offline (available=False, battery_2_offline event, so the shortfall
    #   must lean harder on battery_1 + import).
    # - tick 3: storm_alert=True (a live volatility signal, reduced transmission_headroom too).
    # - ticks 10-12 (evening shortfall) are UNCHANGED from the original 12-row file.
    dict(tick=1, seed=100, objective="", solar_output_mw=0.0, solar_forecast_mw=0.0, wind_output_mw=45.0, wind_forecast_mw=48.0, base_demand_mw=75.0, base_demand_forecast_mw=78.0, industrial_demand_mw=35.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=45.0, battery_1_available=True, battery_2_soc_pct=50.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=42.0, carbon_price_per_ton=28.0, grid_carbon_intensity_t_per_mwh=0.55, demand_response_incentive_per_mwh=18.0, weather_forecast="clear night", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=2, seed=101, objective="", solar_output_mw=0.0, solar_forecast_mw=0.0, wind_output_mw=38.0, wind_forecast_mw=40.0, base_demand_mw=70.0, base_demand_forecast_mw=72.0, industrial_demand_mw=33.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=42.0, battery_1_available=True, battery_2_soc_pct=48.0, battery_2_available=False, previous_floor_pct=25.0, electricity_price_per_mwh=40.0, carbon_price_per_ton=28.0, grid_carbon_intensity_t_per_mwh=0.55, demand_response_incentive_per_mwh=18.0, weather_forecast="pre-dawn, battery_2 under maintenance", storm_alert=False, maintenance_scheduled=True, events="battery_2_offline"),
    dict(tick=3, seed=102, objective="", solar_output_mw=0.0, solar_forecast_mw=0.0, wind_output_mw=20.0, wind_forecast_mw=15.0, base_demand_mw=80.0, base_demand_forecast_mw=85.0, industrial_demand_mw=36.0, grid_frequency_hz=49.95, transmission_constraint_mw=180.0, transmission_headroom_mw=150.0, battery_1_soc_pct=55.0, battery_1_available=True, battery_2_soc_pct=55.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=55.0, carbon_price_per_ton=29.0, grid_carbon_intensity_t_per_mwh=0.5, demand_response_incentive_per_mwh=19.0, weather_forecast="storm approaching, turbines throttled", storm_alert=True, maintenance_scheduled=False, events="storm_alert"),
    dict(tick=4, seed=103, objective="min_carbon", solar_output_mw=20.0, solar_forecast_mw=22.0, wind_output_mw=28.0, wind_forecast_mw=29.0, base_demand_mw=85.0, base_demand_forecast_mw=88.0, industrial_demand_mw=38.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=50.0, battery_1_available=True, battery_2_soc_pct=52.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=28.0, carbon_price_per_ton=29.0, grid_carbon_intensity_t_per_mwh=0.85, demand_response_incentive_per_mwh=18.0, weather_forecast="morning ramp, cheap but dirty grid", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=5, seed=104, objective="max_profit", solar_output_mw=85.0, solar_forecast_mw=86.0, wind_output_mw=30.0, wind_forecast_mw=31.0, base_demand_mw=55.0, base_demand_forecast_mw=57.0, industrial_demand_mw=25.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=12.0, battery_1_soc_pct=50.0, battery_1_available=True, battery_2_soc_pct=50.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=45.0, carbon_price_per_ton=29.0, grid_carbon_intensity_t_per_mwh=0.5, demand_response_incentive_per_mwh=18.0, weather_forecast="midday surplus, line constrained", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=6, seed=105, objective="", solar_output_mw=100.0, solar_forecast_mw=98.0, wind_output_mw=28.0, wind_forecast_mw=27.0, base_demand_mw=58.0, base_demand_forecast_mw=59.0, industrial_demand_mw=26.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=12.0, battery_1_soc_pct=50.0, battery_1_available=True, battery_2_soc_pct=50.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=38.0, carbon_price_per_ton=29.0, grid_carbon_intensity_t_per_mwh=0.5, demand_response_incentive_per_mwh=18.0, weather_forecast="bright midday, biggest surplus of the day", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=7, seed=106, objective="max_renewable_utilisation", solar_output_mw=105.0, solar_forecast_mw=103.0, wind_output_mw=22.0, wind_forecast_mw=21.0, base_demand_mw=60.0, base_demand_forecast_mw=61.0, industrial_demand_mw=27.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=12.0, battery_1_soc_pct=50.0, battery_1_available=True, battery_2_soc_pct=50.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=36.0, carbon_price_per_ton=29.0, grid_carbon_intensity_t_per_mwh=0.48, demand_response_incentive_per_mwh=18.0, weather_forecast="early afternoon surplus, line still constrained", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=8, seed=107, objective="max_profit", solar_output_mw=90.0, solar_forecast_mw=88.0, wind_output_mw=18.0, wind_forecast_mw=18.0, base_demand_mw=100.0, base_demand_forecast_mw=102.0, industrial_demand_mw=40.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=85.0, battery_1_available=True, battery_2_soc_pct=80.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=165.0, carbon_price_per_ton=29.0, grid_carbon_intensity_t_per_mwh=0.45, demand_response_incentive_per_mwh=20.0, weather_forecast="afternoon price spike", storm_alert=False, maintenance_scheduled=False, events="price_spike"),
    dict(tick=9, seed=108, objective="cost_efficiency", solar_output_mw=50.0, solar_forecast_mw=48.0, wind_output_mw=20.0, wind_forecast_mw=19.0, base_demand_mw=115.0, base_demand_forecast_mw=118.0, industrial_demand_mw=42.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=65.0, battery_1_available=True, battery_2_soc_pct=60.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=190.0, carbon_price_per_ton=31.0, grid_carbon_intensity_t_per_mwh=0.32, demand_response_incentive_per_mwh=19.0, weather_forecast="expensive but clean grid", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=10, seed=109, objective="", solar_output_mw=5.0, solar_forecast_mw=2.0, wind_output_mw=32.0, wind_forecast_mw=34.0, base_demand_mw=130.0, base_demand_forecast_mw=132.0, industrial_demand_mw=44.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=55.0, battery_1_available=True, battery_2_soc_pct=50.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=72.0, carbon_price_per_ton=31.0, grid_carbon_intensity_t_per_mwh=0.6, demand_response_incentive_per_mwh=19.0, weather_forecast="dusk, evening ramp-up", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=11, seed=110, objective="", solar_output_mw=0.0, solar_forecast_mw=0.0, wind_output_mw=35.0, wind_forecast_mw=36.0, base_demand_mw=135.0, base_demand_forecast_mw=138.0, industrial_demand_mw=42.0, grid_frequency_hz=49.98, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=38.0, battery_1_available=True, battery_2_soc_pct=35.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=68.0, carbon_price_per_ton=31.0, grid_carbon_intensity_t_per_mwh=0.6, demand_response_incentive_per_mwh=20.0, weather_forecast="evening peak, shortfall", storm_alert=False, maintenance_scheduled=False, events=""),
    dict(tick=12, seed=111, objective="", solar_output_mw=0.0, solar_forecast_mw=0.0, wind_output_mw=38.0, wind_forecast_mw=40.0, base_demand_mw=110.0, base_demand_forecast_mw=108.0, industrial_demand_mw=36.0, grid_frequency_hz=50.0, transmission_constraint_mw=180.0, transmission_headroom_mw=180.0, battery_1_soc_pct=30.0, battery_1_available=True, battery_2_soc_pct=28.0, battery_2_available=True, previous_floor_pct=25.0, electricity_price_per_mwh=52.0, carbon_price_per_ton=30.0, grid_carbon_intensity_t_per_mwh=0.58, demand_response_incentive_per_mwh=18.0, weather_forecast="late night, demand easing", storm_alert=False, maintenance_scheduled=False, events=""),
]


def make_example_file() -> bytes:
    wb = openpyxl.Workbook()
    _write_rows_sheet(wb, EXAMPLE_DAY_ROWS)
    _write_dictionary_sheet(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
